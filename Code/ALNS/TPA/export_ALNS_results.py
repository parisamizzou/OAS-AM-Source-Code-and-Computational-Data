#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Created on Sun Oct  6 14:09:07 2024

@author: parisa
"""


from io import BytesIO

import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
from matplotlib.ticker import MaxNLocator
from openpyxl import load_workbook
from openpyxl.drawing.image import Image



def export_ALNS_results(
    seed,
    running_time,
    initial_solution,
    initial_objective,
    best_solution,
    best_objective,
    best_production_cost,
    best_lateness_cost,
    best_rejection_cost,
    best_unused_capacity_cost,
    best_job_completion_time,
    best_batch_completion_time,
    best_batch_processing_time,
    best_batch_weight,
    best_batch_volume,
    best_batch_height,
    best_batch_unused_capacity,
    best_job_tardiness,
    history_best,
    history_current,
    history_temp,
    history_q,
    destroy_select_count,
    repair_select_count,
    destroy_accepted_count,
    repair_accepted_count,
    destroy_weights,
    repair_weights,
    filename="ALNS_results.xlsx"):
    
    if not (
        len(history_best)
        == len(history_current)
        == len(history_temp)
        == len(history_q)
    ):
        raise ValueError("All ALNS history lists must have the same length.")

    # Identify rejected jobs and active batches.
    total_rejected_jobs = 0
    active_batches = set()
    
    for job, assignment in best_solution.items():
        batch, machine = assignment

        if batch is None and machine is None:
            total_rejected_jobs += 1
        elif batch is None or machine is None:
            raise ValueError(f"Job {job} has an incomplete assignment.")
        else:
            active_batches.add((machine, batch))

    active_batches = sorted(active_batches)

    total_jobs = len(best_solution)
    total_accepted_jobs = total_jobs - total_rejected_jobs
    acceptance_rate = (
        100 * total_accepted_jobs / total_jobs
        if total_jobs > 0
        else 0.0
    )

    # Destroy operator statistics.
    usage_rows = []
    
    destroy_names = sorted(
        set(destroy_select_count)
        | set(destroy_accepted_count)
        | set(destroy_weights)
    )

    for name in destroy_names:
        selected = destroy_select_count.get(name, 0)
        accepted = destroy_accepted_count.get(name, 0)
        rate = 100 * accepted / selected if selected > 0 else 0.0

        usage_rows.append({
            "Operator": name,
            "Type": "Destroy",
            "Selected": selected,
            "Accepted": accepted,
            "Acceptance Rate (%)": rate,
        })

    # Repair operator statistics.
    repair_names = sorted(
        set(repair_select_count)
        | set(repair_accepted_count)
        | set(repair_weights)
    )

    for name in repair_names:
        selected = repair_select_count.get(name, 0)
        accepted = repair_accepted_count.get(name, 0)
        rate = 100 * accepted / selected if selected > 0 else 0.0

        usage_rows.append({
            "Operator": name,
            "Type": "Repair",
            "Selected": selected,
            "Accepted": accepted,
            "Acceptance Rate (%)": rate,
        })

    # Final destroy selection probabilities.
    probability_rows = []
    
    for name, probability in sorted(destroy_weights.items()):
        probability_rows.append({
            "Operator": name,
            "Type": "Destroy",
            "Selection Probability": probability,
        })

    # Final repair selection probabilities.
    for name, probability in sorted(repair_weights.items()):
        probability_rows.append({
            "Operator": name,
            "Type": "Repair",
            "Selection Probability": probability,
        })

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        
        # Seed.
        pd.DataFrame([{"Seed": seed}]).to_excel(writer, sheet_name="Seed", index=False)

        # Initial solution.
        initial_solution_df = pd.DataFrame.from_dict(
            initial_solution,
            orient="index",
            columns=["Batch", "Machine"],)
        
        initial_solution_df.index.name = "Job"
        initial_solution_df.to_excel(writer, sheet_name="Initial_Solution", index=True)

        # Initial objective.
        pd.DataFrame([{"Initial Objective": initial_objective}]).to_excel(writer, sheet_name="Initial_Objective", index=False)
        
        # Best solution.
        best_solution_df = pd.DataFrame.from_dict(best_solution, orient="index", columns=["Batch", "Machine"],)
        best_solution_df.index.name = "Job"
        best_solution_df.to_excel(writer, sheet_name="Best_Solution", index=True)

        # Best objective.
        pd.DataFrame([{"Best Objective": best_objective}]).to_excel(writer, sheet_name="Best_Objective", index=False)

        # Production cost.
        pd.DataFrame([{"Best Production Cost": best_production_cost}]).to_excel(writer, sheet_name="Best_Production_Cost", index=False)
        
        # Lateness cost.
        pd.DataFrame([{"Best Lateness Cost": best_lateness_cost}]).to_excel(writer, sheet_name="Best_Lateness_Cost", index=False)

        # Rejection cost.
        pd.DataFrame([{"Best Rejection Cost": best_rejection_cost}]).to_excel(writer, sheet_name="Best_Rejection_Cost", index=False)

        # Unused-capacity cost.
        pd.DataFrame([{"Best Unused Capacity Cost": best_unused_capacity_cost}]).to_excel(writer, sheet_name="Best_Unused_Capacity_Cost", index=False)

        # Rejected-job count.
        pd.DataFrame([{"Total Rejected Jobs": total_rejected_jobs}]).to_excel(writer, sheet_name="Total_Rejected_Jobs", index=False)
        
        # Acceptance rate.
        pd.DataFrame([{"Acceptance Rate (%)": acceptance_rate}]).to_excel(writer, sheet_name="Acceptance_Rate", index=False)

        # Job completion times.
        job_completion_df = pd.DataFrame.from_dict(best_job_completion_time,orient="index",columns=["Completion Time"],)
        job_completion_df.index.name = "Job"
        job_completion_df.to_excel(writer, sheet_name="Best_Job_Completion_Time", index=True)

        # Job tardiness.
        job_tardiness_df = pd.DataFrame.from_dict(best_job_tardiness,orient="index",columns=["Tardiness"],)
        job_tardiness_df.index.name = "Job"
        job_tardiness_df.to_excel(writer, sheet_name="Best_Job_Tardiness", index=True)
        
        # Batch completion times.
        batch_completion_data = []
        for machine, batch in active_batches:batch_completion_data.append({"Machine": machine,"Batch": batch,"Completion Time":best_batch_completion_time[machine][batch - 1],})
        pd.DataFrame(batch_completion_data,columns=["Machine", "Batch", "Completion Time"],).to_excel(writer, sheet_name="Best_Batch_Completion_Time", index=False)

        # Batch processing times.
        batch_processing_data = []
        for machine, batch in active_batches:
            batch_processing_data.append({
                "Machine": machine,
                "Batch": batch,
                "Processing Time":
                    best_batch_processing_time[machine][batch - 1],})

        pd.DataFrame(
            batch_processing_data,
            columns=["Machine", "Batch", "Processing Time"],).to_excel(
            writer, sheet_name="Best_Batch_Processing_Time", index=False)

        # Batch raw-material weights.
        batch_weight_data = []
        
        for machine, batch in active_batches:
            batch_weight_data.append({
                "Machine": machine,
                "Batch": batch,
                "Weight": best_batch_weight[machine][batch - 1],
            })

        pd.DataFrame(
            batch_weight_data,
            columns=["Machine", "Batch", "Weight"],
        ).to_excel(
            writer, sheet_name="Best_Batch_Weight", index=False
        )

        # Batch actual part volumes.
        batch_volume_data = []
        
        for machine, batch in active_batches:
            batch_volume_data.append({
                "Machine": machine,
                "Batch": batch,
                "Volume": best_batch_volume[machine][batch - 1],})

        pd.DataFrame(
            batch_volume_data,
            columns=["Machine", "Batch", "Volume"],
        ).to_excel(
            writer, sheet_name="Best_Batch_Volume", index=False
        )

        # Batch heights.
        batch_height_data = []

        for machine, batch in active_batches:
            batch_height_data.append({
                "Machine": machine,
                "Batch": batch,
                "Height": best_batch_height[machine][batch - 1],
            })

        pd.DataFrame(
            batch_height_data,
            columns=["Machine", "Batch", "Height"],
        ).to_excel(
            writer, sheet_name="Best_Batch_Height", index=False
        )

        # Include zero unused capacity for full active batches.
        batch_unused_capacity_data = []
        
        for machine, batch in active_batches:
            batch_unused_capacity_data.append({
                "Machine": machine,
                "Batch": batch,
                "Unused Capacity Fraction":
                    best_batch_unused_capacity[machine][batch - 1],
            })

        pd.DataFrame(
            batch_unused_capacity_data,
            columns=["Machine", "Batch", "Unused Capacity Fraction"],
        ).to_excel(
            writer, sheet_name="Best_Batch_Unused_Capacity", index=False
        )

        # Batch utilization.
        batch_utilization_data = []

        for machine, batch in active_batches:
            unused_fraction = best_batch_unused_capacity[machine][batch - 1]

            batch_utilization_data.append({
                "Machine": machine,
                "Batch": batch,
                "Batch Utilization (%)": 100 * (1 - unused_fraction),
            })

        pd.DataFrame(
            batch_utilization_data,
            columns=["Machine", "Batch", "Batch Utilization (%)"],
        ).to_excel(
            writer, sheet_name="Best_Batch_Utilization", index=False
        )

        # Operator usage.
        pd.DataFrame(usage_rows,columns=[
                "Operator", "Type", "Selected", "Accepted",
                "Acceptance Rate (%)",],).to_excel(writer, sheet_name="Operator_Usage", index=False)

        # Final AOS probabilities.
        pd.DataFrame(
            probability_rows,
            columns=["Operator", "Type", "Selection Probability"],
        ).to_excel(
            writer, sheet_name="Final_Probabilities", index=False
        )

        # Best and current objective histories.
        pd.DataFrame({
            "Best Objective": history_best,
            "Current Objective": history_current,
        }).to_excel(
            writer,
            sheet_name="History_Objectives",
            index=True,
            index_label="Iteration",
        )

        # Temperature history.
        pd.DataFrame({
            "Temperature": history_temp,
        }).to_excel(
            writer,
            sheet_name="History_Temperature",
            index=True,
            index_label="Iteration",
        )

        # Destruction-size history.
        pd.DataFrame({
            "q (destruction fraction)": history_q,
        }).to_excel(
            writer,
            sheet_name="History_q",
            index=True,
            index_label="Iteration",
        )

    # Embed charts using in-memory images.
    book = load_workbook(filename)
    image_buffers = []
    
    try:
        graph_sheet = book.create_sheet("Graph_Images")

        # Convergence plot.
        objective_image = BytesIO()
        image_buffers.append(objective_image)

        fig, ax = plt.subplots(figsize=(9, 5))
        try:
            ax.plot(
                history_best,
                label="Best Objective",
                linewidth=2,
                color="red",
            )
            ax.plot(
                history_current,
                label="Current Objective",
                color="steelblue",
                alpha=0.6,
            )
            ax.set_xlabel("Iteration")
            ax.set_ylabel("Objective Value")
            ax.set_title(f"ALNS Convergence (seed={seed})")
            ax.legend()
            ax.grid(True, linestyle="--", alpha=0.4)
            fig.tight_layout()
            fig.savefig(objective_image, format="png", dpi=100)
        finally:
            plt.close(fig)

        objective_image.seek(0)
        graph_sheet.add_image(Image(objective_image), "A1")

        # Operator usage plot.
        operator_image = BytesIO()
        image_buffers.append(operator_image)

        positions = list(range(len(usage_rows)))
        operator_names = [row["Operator"] for row in usage_rows]
        selected_counts = [row["Selected"] for row in usage_rows]
        accepted_counts = [row["Accepted"] for row in usage_rows]
        colors = [
            "#4C72B0" if row["Type"] == "Destroy" else "#DD8452"
            for row in usage_rows
        ]

        fig, ax = plt.subplots(
            figsize=(max(8, len(usage_rows) * 0.7), 5)
        )
        try:
            ax.bar(
                [i - 0.2 for i in positions],
                selected_counts,
                width=0.35,
                color=colors,
                alpha=0.8,
            )
            ax.bar(
                [i + 0.2 for i in positions],
                accepted_counts,
                width=0.35,
                color=colors,
                alpha=0.4,
                hatch="//",
            )
            ax.set_xticks(positions)
            ax.set_xticklabels(
                operator_names, rotation=30, ha="right"
            )
            ax.yaxis.set_major_locator(MaxNLocator(integer=True))
            ax.set_ylabel("Count")
            ax.set_title(f"Operator Usage (seed={seed})")
            ax.legend(handles=[
                Patch(facecolor="#4C72B0", label="Destroy"),
                Patch(facecolor="#DD8452", label="Repair"),
                Patch(facecolor="grey", alpha=0.8, label="Selected"),
                Patch(
                    facecolor="grey",
                    alpha=0.4,
                    hatch="//",
                    label="Accepted",
                ),
            ], fontsize=8)
            ax.grid(axis="y", linestyle="--", alpha=0.4)
            fig.tight_layout()
            fig.savefig(operator_image, format="png", dpi=100)
        finally:
            plt.close(fig)

        operator_image.seek(0)
        graph_sheet.add_image(Image(operator_image), "A32")

        # Temperature plot.
        temperature_image = BytesIO()
        image_buffers.append(temperature_image)

        fig, ax = plt.subplots(figsize=(9, 4))
        try:
            ax.plot(history_temp, color="orange")
            ax.set_xlabel("Iteration")
            ax.set_ylabel("Temperature")
            ax.set_title(f"ALNS Temperature (seed={seed})")
            ax.grid(True, linestyle="--", alpha=0.4)
            fig.tight_layout()
            fig.savefig(temperature_image, format="png", dpi=100)
        finally:
            plt.close(fig)

        temperature_image.seek(0)
        graph_sheet.add_image(Image(temperature_image), "A63")

        # Destruction-size plot.
        destruction_image = BytesIO()
        image_buffers.append(destruction_image)

        fig, ax = plt.subplots(figsize=(9, 4))
        try:
            ax.plot(history_q, color="teal")
            ax.set_xlabel("Iteration")
            ax.set_ylabel("Destruction Fraction q")
            ax.set_title(f"ALNS Destruction Size (seed={seed})")
            ax.grid(True, linestyle="--", alpha=0.4)
            fig.tight_layout()
            fig.savefig(destruction_image, format="png", dpi=100)
        finally:
            plt.close(fig)

        destruction_image.seek(0)
        graph_sheet.add_image(Image(destruction_image), "A89")

        book.save(filename)

    finally:
        book.close()

        for buffer in image_buffers:
            buffer.close()

    print(f"ALNS results successfully exported to {filename}")
        
    
    
    
    
    
   
   

